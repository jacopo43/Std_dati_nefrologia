from __future__ import annotations

import json
import re
import shutil
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import numpy as np
import pandas as pd
from openpyxl import load_workbook

SUPPORTED_EXTENSIONS = {'.xls', '.xlsx', '.ods'}

# These expressions come from other_helps.xlsx.  The workbook is deliberately
# NOT read at run time: it is only the reference used to encode these rules.
NEGATION_PHRASES = (
    'no', 'non', 'nessun', 'nessuna', 'nega', 'assenza di', 'senza',
    'senza evidenza di', 'escluso', 'negative for',
)
FAMILY_PHRASES = (
    'familiarita per', 'anamnesi familiare', 'madre con', 'padre con',
    'family history of',
)
UNCERTAIN_PHRASES = (
    'sospetto', 'possibile', 'probabile', 'da escludere', 'rule out', 'query',
)

# More specific categories prevail over broad/residual categories.
COMORBIDITY_DOMINANCE = {
    'FIBRILLAZIONE ATRIALE': {'ARITMIA'},
    'INSUFFICIENZA CARDIACA': {'SCOMPENSO CARDIACO'},
    'DISTURBI DEL MOVIMENTO': {'DISTURBO NEUROLOGICO'},
    'DEMENZA/ALZHEIMER': {'DISTURBO NEUROLOGICO'},
    'DEPRESSIONE': {'DISTURBI DELL\'UMORE'},
}

# Lower number = higher priority.  Specific clinical consequences precede
# exposure/risk/generic outcomes, as requested in the matching notes.
INTERACTION_PRIORITY = {
    'EMORRAGIA E SANGUINAMENTO': 10,
    'PROLUNGAMENTO QT': 10,
    'IPOTENSIONE': 10,
    'SINCOPE': 10,
    'DEPRESSIONE SNC': 10,
    'MIELOTOSSICITÀ/CITOPENIA': 10,
    'NEFROTOSSICITÀ': 10,
    'EPATOTOSSICITÀ': 10,
    'MIOPATIA/DANNO MUSCOLARE': 10,
    'ALTERAZIONE ELETTROLITICA': 10,
    'IPERGLICEMIA': 10,
    'IPOGLICEMIA': 10,
    'AUMENTO TRANSAMINASI': 20,
    'AUMENTO RISCHIO DI SANGUINAMENTO': 30,
    'AUMENTO RISCHIO CADUTE': 30,
    'RISCHIO FRATTURE': 30,
    'DIMINUZIONE EFFICACIA': 40,
    'ALTERAZIONE ASSORBIMENTO': 45,
    'AUMENTO DELL\'ESPOSIZIONE O DELL\'EFFETTO DI UNO DEI FARMACI': 50,
    'AUMENTO DELLA TOSSICITÀ DI UNO DEI FARMACI': 90,
    'AUMENTO RISCHIO CV': 90,
    'AUMENTO RISCHIO DI EVENTI AVVERSI GASTROINTESTINALI': 90,
    'ALTRO ESITO CLINICAMENTE RILEVANTE': 100,
}


def get_engine(path: Path | str) -> str:
    suffix = Path(path).suffix.lower()
    engines = {'.xlsx': 'openpyxl', '.xls': 'xlrd', '.ods': 'odf'}
    if suffix not in engines:
        raise ValueError(f'Unsupported workbook extension: {suffix}')
    return engines[suffix]


def clean_string(value):
    if value is None or pd.isna(value):
        return None
    value = str(value).strip()
    return value or None


def normalize_text(value) -> str:
    value = clean_string(value)
    if value is None:
        return ''
    value = unicodedata.normalize('NFKD', value.lower())
    value = ''.join(c for c in value if not unicodedata.combining(c))
    value = value.replace('’', "'")
    value = re.sub(r'[^a-z0-9]+', ' ', value)
    return re.sub(r'\s+', ' ', value).strip()


def normalize_name(value) -> str:
    return normalize_text(value).upper()


def split_terms(value, separators=r'[;,]') -> list[str]:
    text = clean_string(value)
    if not text:
        return []
    return [normalize_text(x) for x in re.split(separators, text) if normalize_text(x)]


def boundary_contains(text: str, term: str) -> bool:
    return bool(term and re.search(rf'(?<!\w){re.escape(term)}(?!\w)', text))


def make_unique_columns(columns: Iterable) -> list[str]:
    seen = {}
    out = []
    for value in columns:
        name = clean_string(value) or 'Unnamed'
        seen[name] = seen.get(name, 0) + 1
        out.append(name if seen[name] == 1 else f'{name}__{seen[name]}')
    return out


def read_excel(path: Path | str, **kwargs):
    return pd.read_excel(path, engine=get_engine(path), **kwargs)


def preceding_context(text: str, start: int, max_words: int = 7) -> str:
    prefix = text[:start].strip().split()
    return ' '.join(prefix[-max_words:])


def match_is_blocked(text: str, term: str) -> bool:
    """Exclude negated, family-history, or uncertain mentions."""
    for match in re.finditer(rf'(?<!\w){re.escape(term)}(?!\w)', text):
        context = preceding_context(text, match.start(), max_words=7)
        if any(boundary_contains(context, phrase) for phrase in NEGATION_PHRASES):
            continue
        if any(phrase in context for phrase in FAMILY_PHRASES):
            continue
        if any(boundary_contains(context, phrase) for phrase in UNCERTAIN_PHRASES):
            continue
        return False
    return True


def safe_terms(value, separators=r';') -> set[str]:
    return {x for x in split_terms(value, separators=separators) if len(x) >= 2}


@dataclass(frozen=True)
class DrugInfo:
    substance: str
    atc: str | None


@dataclass(frozen=True)
class DictionaryEntry:
    canonical: str
    automatic_terms: tuple[str, ...]
    roots: tuple[str, ...]
    blocked_terms: tuple[str, ...]
    order: int


class DictionaryCache:
    """Load the current dictionaries once and build indexes for matching."""

    def __init__(self, drugs_path: Path, com_path: Path, interact_path: Path):
        self.drug_aliases: dict[str, DrugInfo] = {}
        self.drug_candidates: dict[str, list[str]] = {}
        self.com_entries: list[DictionaryEntry] = []
        self.interaction_entries: list[DictionaryEntry] = []
        self._load_drugs(drugs_path)
        self._load_comorbidities(com_path)
        self._load_interactions(interact_path)

    def _load_drugs(self, path: Path):
        df = read_excel(path, sheet_name=0).fillna('')
        required = {'canonical_substance', 'aliases'}
        if not required.issubset(df.columns):
            raise ValueError(f'Drugs_dict.xlsx must contain {sorted(required)}')

        canonical_terms = []
        alias_terms = []
        for _, row in df.iterrows():
            canonical = clean_string(row['canonical_substance'])
            if not canonical:
                continue
            info = DrugInfo(canonical, clean_string(row.get('atc_code')))
            canonical_term = normalize_text(canonical)
            if canonical_term:
                canonical_terms.append((canonical_term, info))
            terms = safe_terms(row.get('aliases')) | safe_terms(row.get('match_roots'))
            alias_terms.extend((term, info) for term in terms)

        # Exact canonical names override aliases from combination products.
        for term, info in alias_terms:
            self.drug_aliases.setdefault(term, info)
        for term, info in canonical_terms:
            self.drug_aliases[term] = info

        for alias in self.drug_aliases:
            for token in set(alias.split()):
                if len(token) >= 3:
                    self.drug_candidates.setdefault(token, []).append(alias)
        for token in self.drug_candidates:
            self.drug_candidates[token].sort(key=len, reverse=True)

    @staticmethod
    def _automatic_terms(row, aliases_col, roots_col, usage_col,
                         review_col, non_auto_col) -> tuple[tuple[str, ...], tuple[str, ...], tuple[str, ...]]:
        usage = normalize_name(row.get(usage_col, ''))
        aliases = safe_terms(row.get(aliases_col))
        roots = safe_terms(row.get(roots_col))
        review = safe_terms(row.get(review_col))
        non_auto = safe_terms(row.get(non_auto_col))
        blocked = review | non_auto

        # NO and REVISIONE-only rows must never be assigned automatically.
        automatic_allowed = usage.startswith('SI')
        if not automatic_allowed:
            return (), (), tuple(sorted(aliases | roots | blocked, key=len, reverse=True))

        aliases -= blocked
        roots -= blocked
        return (
            tuple(sorted(aliases, key=len, reverse=True)),
            tuple(sorted(roots, key=len, reverse=True)),
            tuple(sorted(blocked, key=len, reverse=True)),
        )

    def _load_comorbidities(self, path: Path):
        df = read_excel(path, sheet_name=0).fillna('')
        required = {'canonical_name', 'aliases', 'match_roots'}
        if not required.issubset(df.columns):
            raise ValueError(f'Com_dict.xlsx must contain {sorted(required)}')

        for order, (_, row) in enumerate(df.iterrows()):
            canonical = clean_string(row['canonical_name'])
            if not canonical:
                continue
            aliases, roots, blocked = self._automatic_terms(
                row, 'aliases', 'match_roots', 'uso_automatico',
                'aliases_da_revisionare', 'aliases_non_automatici',
            )
            self.com_entries.append(DictionaryEntry(canonical, aliases, roots, blocked, order))

    def _load_interactions(self, path: Path):
        df = read_excel(path, sheet_name=0).fillna('')
        if not {'Conseguenze', 'Effetti'}.issubset(df.columns):
            raise ValueError("Interact_dict.xlsx must contain 'Conseguenze' and 'Effetti'")

        for order, (_, row) in enumerate(df.iterrows()):
            canonical = clean_string(row['Conseguenze'])
            if not canonical:
                continue
            usage = normalize_name(row.get('Uso automatico', ''))
            effects_set = safe_terms(row.get('Effetti'), separators=r',')
            review = safe_terms(row.get('Termini da revisionare'), separators=r',')
            non_auto = safe_terms(row.get('Termini non automatici'), separators=r',')
            blocked_set = review | non_auto
            effects_set -= blocked_set
            effects = tuple(sorted(effects_set, key=len, reverse=True)) if usage.startswith('SI') else ()
            blocked = tuple(sorted(blocked_set, key=len, reverse=True))
            self.interaction_entries.append(DictionaryEntry(canonical, effects, (), blocked, order))

    def lookup_drug(self, text) -> DrugInfo | None:
        normalized = normalize_text(text)
        if not normalized:
            return None
        if normalized in self.drug_aliases:
            return self.drug_aliases[normalized]

        candidates = set()
        for token in normalized.split():
            candidates.update(self.drug_candidates.get(token, ()))
        for alias in sorted(candidates, key=len, reverse=True):
            if boundary_contains(normalized, alias):
                return self.drug_aliases[alias]
        return None

    def find_drugs(self, text) -> list[DrugInfo]:
        normalized = normalize_text(text)
        if not normalized:
            return []
        candidates = set()
        for token in normalized.split():
            candidates.update(self.drug_candidates.get(token, ()))
        found = {}
        for alias in sorted(candidates, key=len, reverse=True):
            if boundary_contains(normalized, alias):
                info = self.drug_aliases[alias]
                found.setdefault(normalize_text(info.substance), info)
        return list(found.values())

    @staticmethod
    def _entry_matches(normalized: str, entry: DictionaryEntry) -> bool:
        for term in entry.automatic_terms:
            if boundary_contains(normalized, term) and not match_is_blocked(normalized, term):
                return True
        for root in entry.roots:
            # Roots can be stems, therefore boundary matching is not required.
            for match in re.finditer(re.escape(root), normalized):
                context = preceding_context(normalized, match.start(), max_words=7)
                if any(boundary_contains(context, x) for x in NEGATION_PHRASES):
                    continue
                if any(x in context for x in FAMILY_PHRASES):
                    continue
                if any(boundary_contains(context, x) for x in UNCERTAIN_PHRASES):
                    continue
                return True
        return False

    def find_comorbidities(self, text) -> list[str]:
        normalized = normalize_text(text)
        if not normalized:
            return []

        found = [entry.canonical for entry in self.com_entries if self._entry_matches(normalized, entry)]

        # The current reference groups all infection sites into one category.
        infection_matches = [x for x in found if x.startswith('INFEZIONE') or 'MALATTIA INFETTIVA' in x]
        if infection_matches:
            preferred = next((x for x in found if x == 'INFEZIONE/MALATTIA INFETTIVA'), infection_matches[0])
            found = [x for x in found if x not in infection_matches] + [preferred]

        # Apply specific-over-generic collision rules.
        found_set = set(found)
        for specific, generic_set in COMORBIDITY_DOMINANCE.items():
            if specific in found_set:
                found_set -= generic_set

        # Residual ALTRO is never automatic.
        found_set.discard('ALTRO')
        output = []
        for value in found:
            if value in found_set and value not in output:
                output.append(value)
        return output

    def has_comorbidity(self, text, target: str) -> bool:
        return normalize_text(target) in {normalize_text(x) for x in self.find_comorbidities(text)}

    def find_consequences(self, text, max_items=3) -> list[str]:
        normalized = normalize_text(text)
        if not normalized:
            return []

        matches = [entry for entry in self.interaction_entries if self._entry_matches(normalized, entry)]
        names = {x.canonical for x in matches}

        # Small grammatical heuristics cover common Italian inflections that
        # cannot be represented reliably by exact comma-separated aliases.
        if re.search(r'\b(ast|alt|got|gpt)\b.{0,25}\b(aumentat|elevat)', normalized) or re.search(
            r'\b(aumentat|elevat)\b.{0,25}\b(ast|alt|got|gpt)\b', normalized
        ):
            names.add('AUMENTO TRANSAMINASI')
        if re.search(r'\b(aumenta|aumentano|incrementa|incrementano)\b.{0,35}\b(concentrazion|livell)', normalized):
            names.add('AUMENTO DELL\'ESPOSIZIONE O DELL\'EFFETTO DI UNO DEI FARMACI')

        risk_bleeding = bool(re.search(r'\brischio\b.{0,30}\b(sanguin|emorrag)', normalized))
        manifest_bleeding = bool(re.search(
            r'\b(emorragia|sanguinamento|melena|ematemesi|ematuria|epistassi|ematoma|ecchimosi)\b',
            normalized,
        )) and not risk_bleeding
        if risk_bleeding:
            names.add('AUMENTO RISCHIO DI SANGUINAMENTO')
            names.discard('EMORRAGIA E SANGUINAMENTO')
        elif manifest_bleeding and 'EMORRAGIA E SANGUINAMENTO' in names:
            names.discard('AUMENTO RISCHIO DI SANGUINAMENTO')

        # Organ damage prevails over laboratory-only change.
        if 'EPATOTOSSICITÀ' in names:
            names.discard('AUMENTO TRANSAMINASI')
        # Specific toxicity prevails over generic exposure/toxicity categories.
        specific = {x for x in names if INTERACTION_PRIORITY.get(x, 50) <= 30}
        if specific:
            names.discard('AUMENTO DELL\'ESPOSIZIONE O DELL\'EFFETTO DI UNO DEI FARMACI')
            names.discard('AUMENTO DELLA TOSSICITÀ DI UNO DEI FARMACI')

        ordered = sorted(
            names,
            key=lambda name: (
                INTERACTION_PRIORITY.get(name, 50),
                next((e.order for e in matches if e.canonical == name), 9999),
            ),
        )
        return ordered[:max_items]


class RuleResult:
    def __init__(self, data=None, status='mapped', message=None):
        self.data = pd.DataFrame() if data is None else data.reset_index(drop=True)
        self.status = status
        self.message = message


class ConverterEngine:
    # The look-ahead accepts compact strings such as '5 mg1 cp' while avoiding
    # partial unit matches inside alphabetic words.
    DOSAGE_PATTERN = re.compile(
        r'(?P<value>\d+(?:[.,]\d+)?)\s*(?P<unit>kg|mg|mcg|ug|µg|g|ui|iu)(?=$|[^a-zA-Zµ])',
        re.I,
    )
    ROUTE_PATTERN = re.compile(r'\b(per\s+os|orale|os|ev|iv|im|sc|sottocute|sottocutanea|inalatoria|topica|transdermica)\b', re.I)

    def __init__(self, dictionaries: DictionaryCache):
        self.d = dictionaries

    def apply(self, rule, values, parameters):
        method = getattr(self, f'rule_{rule}', None)
        if method is None:
            raise ValueError(f'Unknown rule: {rule}')
        return method(values, parameters)

    @staticmethod
    def cleaned(values):
        return [str(v).strip() for v in values if v is not None and not pd.isna(v) and str(v).strip()]

    def rule_parse_numeric(self, values, p):
        vals = self.cleaned(values)
        if not vals:
            return RuleResult(pd.DataFrame({'value': [np.nan]}), 'source_found_but_empty')
        raw = vals[0].replace(',', '.')
        m = re.search(r'-?\d+(?:\.\d+)?', raw)
        value = pd.to_numeric(m.group() if m else raw, errors='coerce')
        return RuleResult(pd.DataFrame({'value': [value]}), 'mapped' if pd.notna(value) else 'numeric_parse_failed')

    def rule_parse_date(self, values, p):
        vals = self.cleaned(values)
        value = (
            pd.to_datetime(
                vals[0],
                dayfirst=p.get('dayfirst', True),
                errors=p.get('errors', 'coerce'),
            )
            if vals
            else pd.NaT
        )
        if pd.isna(value):
            return RuleResult(
                pd.DataFrame({'value': [np.nan]}),
                'date_parse_failed',
            )

        # Dates are intentionally exported as ISO strings rather than Excel
        # datetimes, preventing values such as '2026-05-25 00:00:00'.
        date_format = p.get('format', '%Y-%m-%d')
        return RuleResult(pd.DataFrame({'value': [value.strftime(date_format)]}))

    def rule_calculate_age(self, values, p):
        vals = self.cleaned(values)
        birth = pd.to_datetime(vals[0], dayfirst=p.get('dayfirst', True), errors='coerce') if vals else pd.NaT
        ref = pd.to_datetime(p.get('reference_date', '2026-01-01'), errors='coerce')
        if pd.isna(birth) or pd.isna(ref):
            return RuleResult(pd.DataFrame({'age_years': [np.nan]}), 'age_calculation_failed')
        age = ref.year - birth.year - int((ref.month, ref.day) < (birth.month, birth.day))
        return RuleResult(pd.DataFrame({'age_years': [age]}))

    def rule_count_nonempty_rows(self, values, p):
        return RuleResult(pd.DataFrame({'count': [len(self.cleaned(values))]}))

    def rule_count_drug_mentions(self, values, p):
        unique = {
            normalize_text(info.substance)
            for value in self.cleaned(values)
            for info in self.d.find_drugs(value)
        }
        return RuleResult(pd.DataFrame({'count': [len(unique)]}))

    def rule_count_comorbidities(self, values, p):
        unique = {
            comorbidity
            for value in self.cleaned(values)
            for comorbidity in self.d.find_comorbidities(value)
            if comorbidity not in {'FUMO', 'ALCOL', 'ALCOOL'}
        }
        return RuleResult(pd.DataFrame({'count': [len(unique)]}))

    def rule_row_aligned_value(self, values, p):
        rows = [{'value': clean_string(value), '_source_row': idx} for idx, value in enumerate(values)]
        df = pd.DataFrame(rows)
        if p.get('forward_fill') and not df.empty:
            df['value'] = df['value'].ffill()
        if p.get('drop_empty', True) and not df.empty:
            df = df[df['value'].notna()]
        return RuleResult(df, 'mapped' if not df.empty else 'source_found_but_empty')

    def rule_comorbidity_dictionary_lookup(self, values, p):
        rows = []
        for idx, value in enumerate(values):
            for subidx, comorbidity in enumerate(self.d.find_comorbidities(value)):
                if comorbidity in {'FUMO', 'ALCOL', 'ALCOOL'}:
                    continue
                rows.append({'comorbidity': comorbidity, '_source_row': idx, '_source_subrow': subidx})
        df = pd.DataFrame(rows)
        if not p.get('keep_duplicates', False) and not df.empty:
            df = df.drop_duplicates('comorbidity', keep='first')
        return RuleResult(df, 'mapped' if not df.empty else 'no_comorbidities_extracted')

    def rule_comorbidity_flag_lookup(self, values, p):
        target = p['target']
        found = any(self.d.has_comorbidity(value, target) for value in self.cleaned(values))
        return RuleResult(pd.DataFrame({'value': [1 if found else 0]}))

    def rule_interaction_dictionary_lookup(self, values, p):
        rows = []
        max_items = int(p.get('max_consequences', 3))
        for idx, value in enumerate(values):
            consequences = self.d.find_consequences(value, max_items=max_items)
            row = {'_source_row': idx}
            for j in range(1, max_items + 1):
                row[f'consequence_{j}'] = consequences[j - 1] if j <= len(consequences) else np.nan
            rows.append(row)
        df = pd.DataFrame(rows)
        return RuleResult(df, 'mapped' if not df.empty else 'source_found_but_empty')

    def rule_drug_dictionary_lookup(self, values, p):
        rows = []
        for idx, value in enumerate(values):
            info = self.d.lookup_drug(value)
            rows.append({
                'substance': info.substance if info else np.nan,
                'atc': info.atc if info else np.nan,
                '_source_row': idx,
            })
        df = pd.DataFrame(rows)
        status = 'mapped' if not df.empty and df['substance'].notna().any() else 'dictionary_lookup_failed'
        return RuleResult(df, status)

    def rule_parse_treatment_row(self, values, p):
        rows = []
        for idx, value in enumerate(values):
            text = clean_string(value)
            if not text:
                continue
            dose = self.DOSAGE_PATTERN.search(text)
            route = self.ROUTE_PATTERN.search(text)
            dosage_value = float(dose.group('value').replace(',', '.')) if dose else np.nan
            dosage_unit = dose.group('unit').lower() if dose else np.nan
            daily_dose = self._daily_dose(text, dosage_value)
            rows.append({
                'dosage_value': dosage_value,
                'dosage_unit': dosage_unit,
                'daily_dose': daily_dose,
                'route': route.group().strip() if route else np.nan,
                '_source_row': idx,
            })
        df = pd.DataFrame(rows)
        return RuleResult(df, 'mapped' if not df.empty else 'source_found_but_empty')

    @staticmethod
    def _daily_dose(text, dose):
        """Return administrations per day, not milligrams per day.

        Examples:
        - '5 mg 1 cp ore 8' -> 1
        - '180 mg ore 8 e ore 20' -> 2
        - '2,5 mg (2 cp da 1 mg) ore 8' -> 1
        - '5 mg/24ore cerotto transdermico' -> 1
        """
        if pd.isna(dose):
            return np.nan

        raw = str(text).lower().replace('’', "'")
        normalized = normalize_text(text)

        # Explicit frequency always has priority.
        explicit_patterns = [
            r'(\d+(?:[.,]\d+)?)\s*volte\s*(?:al|a|il)?\s*(?:giorno|die)',
            r'(\d+(?:[.,]\d+)?)\s*(?:x|per)\s*(?:die|giorno)',
        ]
        for pattern in explicit_patterns:
            match = re.search(pattern, raw, flags=re.I)
            if match:
                return float(match.group(1).replace(',', '.'))

        every_hours = re.search(r'ogni\s*(\d+(?:[.,]\d+)?)\s*ore?', raw, flags=re.I)
        if every_hours:
            hours = float(every_hours.group(1).replace(',', '.'))
            return round(24 / hours, 6) if hours > 0 else np.nan

        # A formulation lasting 24 hours represents one daily administration.
        if re.search(r'(?:/|ogni\s*)24\s*ore?', raw, flags=re.I):
            return 1

        # Count distinct administration times. Repeated mentions of the same
        # hour do not inflate the frequency.
        hours = re.findall(r'\bore?\s*(\d{1,2})(?::([0-5]\d))?', raw, flags=re.I)
        if hours:
            unique_hours = {(int(hour), minute or '00') for hour, minute in hours}
            return len(unique_hours)

        # Common textual schedules.
        if any(phrase in normalized for phrase in [
            'tre volte al giorno', 'ter die', 'tid',
            'mattina pomeriggio sera',
        ]):
            return 3
        if any(phrase in normalized for phrase in [
            'due volte al giorno', 'bis die', 'bid',
            'mattina e sera', 'mattino e sera',
        ]):
            return 2
        if any(phrase in normalized for phrase in [
            'una volta al giorno', '1 die', 'qd',
            'al mattino', 'alla sera', 'la sera', 'la mattina',
        ]):
            return 1

        # A single tablet/capsule/patch instruction without another frequency
        # indicator is interpreted as one administration per day. Tablet count
        # is deliberately not used as frequency: '2 cp da 1 mg ore 8' is one dose.
        if re.search(r'\b\d+(?:[.,]\d+)?\s*(?:cp|cpr|compress[ae]|capsul[ae]|cerott[oi])\b', raw, flags=re.I):
            return 1

        return np.nan


def parse_parameters(value):
    if value is None or pd.isna(value) or str(value).strip() == '':
        return {}
    return json.loads(value) if not isinstance(value, dict) else dict(value)


def read_link(link_path):
    link = read_excel(link_path, sheet_name='LINK').fillna('')
    required = [
        'sheet_name', 'group_name', 'old_name', 'new_name', 'short_name',
        'source', 'rule', 'rule_group', 'output_field', 'expansion_group',
        'rule_parameters',
    ]
    missing = [x for x in required if x not in link.columns]
    if missing:
        raise ValueError(f'Missing LINK columns: {missing}')
    link = link[(link['short_name'] != '') & (link['rule'] != '')].copy().reset_index(drop=True)
    link['_link_id'] = [f'__link_{i:04d}' for i in range(len(link))]
    return link


def read_patient_sheets(path):
    book = pd.ExcelFile(path, engine=get_engine(path))
    out = {}
    for sheet in book.sheet_names:
        # Patient workbooks use row 2 as the variable-name row.
        df = read_excel(path, sheet_name=sheet, header=1)
        df.columns = make_unique_columns(df.columns)
        out[sheet] = df
    return out


def find_source(df, source):
    lookup = {normalize_name(c.split('__')[0]): c for c in df.columns}
    col = lookup.get(normalize_name(source))
    return (col, df[col]) if col is not None else (None, None)


def combine_frames(frames):
    if not frames:
        return pd.DataFrame(index=[0])
    keyed, singles = [], []
    for frame in frames:
        frame = frame.copy()
        keys = [x for x in ['_source_row', '_source_subrow'] if x in frame.columns]
        if keys:
            keyed.append(frame.drop_duplicates(keys).set_index(keys))
        elif len(frame) == 1:
            singles.append(frame.reset_index(drop=True))
        else:
            keyed.append(frame)
    result = pd.concat(keyed, axis=1, join='outer').reset_index() if keyed else pd.DataFrame(index=[0])
    result = result.loc[:, ~result.columns.duplicated()]
    n = max(len(result), 1)
    for frame in singles:
        repeated = pd.concat([frame] * n, ignore_index=True)
        result = pd.concat([result.reset_index(drop=True), repeated], axis=1)
    return result.loc[:, ~result.columns.duplicated()]


def execute_patient(patient_id, sheet_name, df, link, engine):
    """Execute LINK rules and return one frame for each expansion group."""
    logs = []
    expansion_frames = {}

    for expansion, expansion_df in link.groupby('expansion_group', sort=False):
        group_outputs = []
        for _, group in expansion_df.groupby('rule_group', sort=False):
            execution_outputs = []
            for (source, rule, params_raw), rows in group.groupby(
                ['source', 'rule', 'rule_parameters'], sort=False, dropna=False
            ):
                source_col, values = find_source(df, source)
                if source_col is None:
                    result = RuleResult(pd.DataFrame(index=[0]), 'source_not_found_in_patient_sheet')
                else:
                    result = engine.apply(rule, values, parse_parameters(params_raw))

                out = pd.DataFrame(index=result.data.index if not result.data.empty else [0])
                for key in ['_source_row', '_source_subrow']:
                    if key in result.data.columns:
                        out[key] = result.data[key].values

                for _, link_row in rows.iterrows():
                    field = link_row['output_field']
                    link_id = link_row['_link_id']
                    out[link_id] = result.data[field].values if field in result.data.columns else np.nan
                    logs.append({
                        'patient_id': patient_id,
                        'patient_sheet_name': sheet_name,
                        'target_sheet': link_row['sheet_name'],
                        'group_name': link_row['group_name'],
                        'old_name': link_row['old_name'],
                        'new_name': link_row['new_name'],
                        'short_name': link_row['short_name'],
                        'source': source,
                        'raw_column_found': source_col,
                        'rule': rule,
                        'status': result.status,
                        'n_extracted_rows': len(result.data),
                        'message': result.message,
                    })
                execution_outputs.append(out)
            group_outputs.append(combine_frames(execution_outputs))
        expansion_frames[expansion] = combine_frames(group_outputs)

    expansion_frames.setdefault('patient', pd.DataFrame(index=[0]))
    expansion_frames['patient'] = expansion_frames['patient'].iloc[[0]].reset_index(drop=True)
    return expansion_frames, logs


@dataclass
class TemplateSheet:
    name: str
    header_rows: int
    groups: list[str]
    old_headers: list[str]
    output_headers: list[str]
    link_ids: list[str | None]


class OutputTemplate:
    """Read output structure from dataset_finale_base.xlsx."""

    def __init__(self, template_path: Path, link: pd.DataFrame):
        self.path = Path(template_path)
        self.sheet_order: list[str] = []
        self.sheets: dict[str, TemplateSheet] = {}
        self._load(link)

    @staticmethod
    def _find_link_row(link, sheet_name, group_name, old_name):
        sheet_candidates = link[
            link['sheet_name'].map(normalize_name) == normalize_name(sheet_name)
        ].copy()
        if group_name:
            # In GENERALE, the same old label occurs at several timepoints.
            # A mapping is valid only for the explicitly named group.
            sheet_candidates = sheet_candidates[
                sheet_candidates['group_name'].map(normalize_name) == normalize_name(group_name)
            ]

        old_norm = normalize_name(old_name)
        exact = sheet_candidates[
            sheet_candidates['old_name'].map(normalize_name) == old_norm
        ]
        if not exact.empty:
            return exact.iloc[0]

        # Permit explanatory suffixes in the template, e.g.
        # 'COMORBIDITA PRE MR (MH: dizionario da fare)'.
        prefixed = sheet_candidates[
            sheet_candidates['old_name'].map(
                lambda x: bool(normalize_name(x)) and old_norm.startswith(normalize_name(x))
            )
        ]
        return None if prefixed.empty else prefixed.iloc[0]

    def _load(self, link):
        wb = load_workbook(self.path, read_only=True, data_only=False)
        for ws in wb.worksheets:
            self.sheet_order.append(ws.title)
            header_rows = 2 if normalize_name(ws.title) == 'GENERALE' else 1
            max_col = ws.max_column
            groups = []
            old_headers = []
            output_headers = []
            link_ids = []

            for col in range(1, max_col + 1):
                group = clean_string(ws.cell(1, col).value) if header_rows == 2 else None
                old_header = clean_string(ws.cell(header_rows, col).value)
                groups.append(group or '')
                old_headers.append(old_header or '')

                matched = self._find_link_row(link, ws.title, group or '', old_header or '')
                if matched is None:
                    output_headers.append(old_header or '')
                    link_ids.append(None)
                else:
                    output_headers.append(clean_string(matched['new_name']) or old_header or '')
                    link_ids.append(matched['_link_id'])

            self.sheets[ws.title] = TemplateSheet(
                name=ws.title,
                header_rows=header_rows,
                groups=groups,
                old_headers=old_headers,
                output_headers=output_headers,
                link_ids=link_ids,
            )
        wb.close()


def _sheet_expansion_groups(sheet_name, link):
    rows = link[link['sheet_name'].map(normalize_name) == normalize_name(sheet_name)]
    groups = [x for x in rows['expansion_group'].tolist() if clean_string(x)]
    return list(dict.fromkeys(groups))


def make_patient_sheet_frame(patient_id, template_sheet, link, expansion_frames):
    relevant_groups = _sheet_expansion_groups(template_sheet.name, link)
    non_patient_groups = [x for x in relevant_groups if x != 'patient']

    if non_patient_groups:
        base = combine_frames([expansion_frames.get(x, pd.DataFrame(index=[0])) for x in non_patient_groups])
    else:
        base = pd.DataFrame(index=[0])

    if base.empty:
        base = pd.DataFrame(index=[0])
    base = base.reset_index(drop=True)
    n = max(len(base), 1)

    # Broadcast patient-level values onto expanded sheets (e.g. DATA RICOGNIZIONE).
    patient_frame = expansion_frames.get('patient', pd.DataFrame(index=[0])).iloc[[0]].reset_index(drop=True)
    if not patient_frame.empty:
        patient_repeated = pd.concat([patient_frame] * n, ignore_index=True)
        base = pd.concat([base, patient_repeated], axis=1)
        base = base.loc[:, ~base.columns.duplicated()]

    output = pd.DataFrame(index=range(n))
    for idx, (old_header, link_id) in enumerate(zip(template_sheet.old_headers, template_sheet.link_ids)):
        internal_col = f'__template_col_{idx:03d}'
        if normalize_name(old_header) == 'ID PAZIENTE':
            output[internal_col] = [patient_id] * n
        elif link_id and link_id in base.columns:
            output[internal_col] = base[link_id].reindex(range(n)).values
        else:
            output[internal_col] = [np.nan] * n

    if normalize_name(template_sheet.name) == 'INTERAZIONI':
        output = clean_interaction_rows(output, template_sheet)

    return output.reset_index(drop=True)



def clean_interaction_rows(output, template_sheet):
    """Remove artifacts created by forward-filled interaction drug columns.

    Drug A and Drug B are deliberately forward-filled because merged cells are
    common in the source workbook.  This can create trailing rows containing
    only the same drug pair while type, reason, consequence, and clinical advice
    are all empty.  Such rows are removed when a populated record for the same
    pair exists.  A pair-only row is retained when it is the sole record for
    that pair, so potentially useful source information is not discarded.
    """
    if output.empty:
        return output

    def matching_columns(*terms):
        cols = []
        for idx, (old_header, new_header) in enumerate(
            zip(template_sheet.old_headers, template_sheet.output_headers)
        ):
            label = f"{old_header} {new_header}"
            normalized = normalize_name(label)
            if any(normalize_name(term) in normalized for term in terms):
                cols.append(f'__template_col_{idx:03d}')
        return cols

    drug_1_cols = matching_columns('FARMACO INTERAGENTE 1', 'FARMACO INTERAGENTE_1')
    drug_2_cols = matching_columns('FARMACO INTERAGENTE 2', 'FARMACO INTERAGENTE_2')
    pair_cols = (drug_1_cols[:1] + drug_2_cols[:1])

    detail_cols = matching_columns(
        'TIPO INTERAZIONE',
        'TIPO DI INTERAZIONE',
        'MOTIVO INTERAZIONE',
        'MOTIVO DELL INTERAZIONE',
        'CONSEGUENZA INTERAZIONE',
        'CONSEGUENZA DELL INTERAZIONE',
        'CONSIGLIO CLINICO',
    )

    if len(pair_cols) < 2:
        return output.drop_duplicates().reset_index(drop=True)

    def populated(value):
        return clean_string(value) is not None

    pair_keys = output[pair_cols].apply(
        lambda row: tuple(normalize_text(value) for value in row), axis=1
    )
    has_pair = output[pair_cols].apply(lambda row: any(populated(v) for v in row), axis=1)
    has_details = (
        output[detail_cols].apply(lambda row: any(populated(v) for v in row), axis=1)
        if detail_cols else pd.Series(False, index=output.index)
    )

    populated_pairs = set(pair_keys[has_pair & has_details])
    inherited_blank = has_pair & ~has_details & pair_keys.isin(populated_pairs)
    cleaned = output.loc[~inherited_blank].copy()

    # Remove fully identical interaction records, while retaining the first
    # occurrence and the original clinical order.
    cleaned = cleaned.drop_duplicates(keep='first')
    return cleaned.reset_index(drop=True)

def excel_value(value):
    if value is None or pd.isna(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if isinstance(value, np.generic):
        return value.item()
    return value


def write_template_output(template_path, output_path, template, sheet_frames):
    """Copy the template, retain its formatting, rename mapped headers, and add data."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template_path, output_path)
    wb = load_workbook(output_path)

    for sheet_name in template.sheet_order:
        ws = wb[sheet_name]
        spec = template.sheets[sheet_name]

        # Remove any old data while preserving the template header/style rows.
        if ws.max_row > spec.header_rows:
            ws.delete_rows(spec.header_rows + 1, ws.max_row - spec.header_rows)

        if spec.header_rows == 2:
            # Row 1 is intentionally retained to distinguish identical labels
            # belonging to different timepoints.
            for col, group in enumerate(spec.groups, 1):
                ws.cell(1, col).value = group or None
            for col, header in enumerate(spec.output_headers, 1):
                ws.cell(2, col).value = header
        else:
            for col, header in enumerate(spec.output_headers, 1):
                ws.cell(1, col).value = header

        frame = sheet_frames.get(sheet_name, pd.DataFrame())
        start_row = spec.header_rows + 1
        for row_offset, row in enumerate(frame.itertuples(index=False, name=None)):
            for col_offset, value in enumerate(row, 1):
                ws.cell(start_row + row_offset, col_offset).value = excel_value(value)

    wb.save(output_path)


def build_all_dataset_from_link(
    link_path,
    raw_patient_path,
    output_path,
    log_path,
    drugs_dict_path,
    com_dict_path,
    interact_dict_path,
    template_path,
):
    link = read_link(link_path)
    template = OutputTemplate(Path(template_path), link)
    dictionaries = DictionaryCache(Path(drugs_dict_path), Path(com_dict_path), Path(interact_dict_path))
    engine = ConverterEngine(dictionaries)
    patients = read_patient_sheets(raw_patient_path)

    all_sheet_frames = {name: [] for name in template.sheet_order}
    logs = []

    for patient_id, (patient_sheet_name, patient_df) in enumerate(patients.items(), 1):
        expansion_frames, patient_logs = execute_patient(
            patient_id, patient_sheet_name, patient_df, link, engine
        )
        logs.extend(patient_logs)

        for sheet_name in template.sheet_order:
            frame = make_patient_sheet_frame(
                patient_id,
                template.sheets[sheet_name],
                link,
                expansion_frames,
            )
            all_sheet_frames[sheet_name].append(frame)

    combined_sheets = {
        name: (pd.concat(frames, ignore_index=True) if frames else pd.DataFrame())
        for name, frames in all_sheet_frames.items()
    }

    write_template_output(
        template_path=template_path,
        output_path=output_path,
        template=template,
        sheet_frames=combined_sheets,
    )

    log_df = pd.DataFrame(logs)
    if log_path:
        log_path = Path(log_path)
        log_path.parent.mkdir(parents=True, exist_ok=True)
        with pd.ExcelWriter(log_path, engine='openpyxl') as writer:
            log_df.to_excel(writer, sheet_name='mapping_log', index=False)
            if not log_df.empty:
                (
                    log_df['status']
                    .value_counts(dropna=False)
                    .rename_axis('status')
                    .reset_index(name='frequency')
                    .to_excel(writer, sheet_name='status_summary', index=False)
                )

    return combined_sheets, log_df


